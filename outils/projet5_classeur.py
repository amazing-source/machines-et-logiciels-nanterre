# -*- coding: utf-8 -*-
# Seance 5, projet : construction du classeur (donnees + formules + synthese).
# Les tris, sous-totaux, filtre avance, TCD, mises en forme conditionnelles et
# graphiques sont ensuite poses dans Excel par projet5_excel.ps1.
# Usage : python projet5_classeur.py <brut.csv> <classeur.xlsx> <identite.json>
import csv
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SRC, OUT = sys.argv[1], sys.argv[2]
ID = json.load(open(sys.argv[3], encoding="utf-8"))

NAVY, ORANGE, INK, MUTED = "12284C", "C8501E", "0B0B0B", "52514E"
PALE, PALE_O, LINE = "EAF1FB", "FCEFE8", "C3C2B7"

rows = list(csv.DictReader(open(SRC, encoding="utf-8-sig"), delimiter=";"))

# --- nettoyage documente des libelles ---------------------------------------
FILIERE = {"Licence_Las": "Licence L.AS", "Ecole d'Ingénieur": "École d'ingénieur",
           "Ecole de Commerce": "École de commerce"}
SELECT = {"formation sélective": "Sélective", "formation non sélective": "Non sélective"}

wb = Workbook()
F = lambda **k: Font(name="Calibri", size=k.pop("size", 11), **k)
thin = Side(style="thin", color=LINE)
box = Border(left=thin, right=thin, top=thin, bottom=thin)


def head(ws, ref, fill=NAVY):
    c = ws[ref]
    c.font = F(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=fill)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = box


# =============================================================================
# Feuille Donnees
# =============================================================================
ws = wb.active
ws.title = "Données"
N = len(rows)
LAST = N + 1

cols = [  # (titre, largeur, calculee ?)
    ("Établissement", 27, False), ("Code dép.", 9, False), ("Département", 16, False),
    ("Filière", 17, False), ("Formation", 34, False), ("Sélectivité", 14, False),
    ("Places", 10, False), ("Candidatures", 15, False), ("Admis", 9, False),
    ("Zone", 16, True), ("Candidatures par place", 15, True),
    ("Taux de remplissage", 15, True), ("Niveau de tension", 13, True),
]
for j, (t, w, calc) in enumerate(cols, 1):
    ws.cell(1, j, t)
    head(ws, "%s1" % get_column_letter(j), ORANGE if calc else NAVY)
    # a gauche : le bouton de filtre occupe le coin inferieur droit de l'en-tete
    ws.cell(1, j).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=0)
    ws.column_dimensions[get_column_letter(j)].width = w
ws.row_dimensions[1].height = 32

for i, r in enumerate(rows, 2):
    ws.cell(i, 1, r["g_ea_lib_vx"])
    ws.cell(i, 2, int(r["dep"]))
    ws.cell(i, 3, r["dep_lib"])
    ws.cell(i, 4, FILIERE.get(r["fili"], r["fili"]))
    ws.cell(i, 5, r["lib_for_voe_ins"])
    ws.cell(i, 6, SELECT[r["select_form"]])
    ws.cell(i, 7, int(r["capa_fin"]))
    ws.cell(i, 8, int(r["voe_tot"]))
    ws.cell(i, 9, int(r["acc_tot"]))
    # colonnes calculees
    ws.cell(i, 10, "=VLOOKUP(B%d,'Paramètres'!$A$4:$C$11,3,FALSE)" % i)
    ws.cell(i, 11, '=IFERROR(H%d/G%d,"")' % (i, i))
    ws.cell(i, 12, '=IFERROR(I%d/G%d,"")' % (i, i))
    ws.cell(i, 13, ('=IF(K%d>=\'Paramètres\'!$F$6,"Très forte",IF(K%d>=\'Paramètres\'!$F$5,"Forte",'
                    'IF(K%d>=\'Paramètres\'!$F$4,"Modérée","Faible")))') % (i, i, i))
    ws.cell(i, 2).alignment = Alignment(horizontal="center")
    for j in (7, 8, 9):
        ws.cell(i, j).number_format = "#,##0"
    ws.cell(i, 11).number_format = "0.0"
    ws.cell(i, 12).number_format = "0%"
    ws.cell(i, 13).alignment = Alignment(horizontal="center")

ws.freeze_panes = "B2"

# =============================================================================
# Feuille Parametres
# =============================================================================
wp = wb.create_sheet("Paramètres")
wp["A1"] = "Tables de référence utilisées par les formules de la feuille Données"
wp["A1"].font = F(bold=True, size=13, color=NAVY)
for ref, t in (("A3", "Code dép."), ("B3", "Département"), ("C3", "Zone")):
    wp[ref] = t
    head(wp, ref)
ZONES = [(75, "Paris", "Paris"),
         (77, "Seine-et-Marne", "Grande couronne"), (78, "Yvelines", "Grande couronne"),
         (91, "Essonne", "Grande couronne"),
         (92, "Hauts-de-Seine", "Petite couronne"), (93, "Seine-Saint-Denis", "Petite couronne"),
         (94, "Val-de-Marne", "Petite couronne"), (95, "Val-d'Oise", "Grande couronne")]
for k, z in enumerate(ZONES, 4):
    for j, v in enumerate(z, 1):
        c = wp.cell(k, j, v)
        c.border = box
    wp.cell(k, 1).alignment = Alignment(horizontal="center")
wp["A13"] = "Utilisation : =RECHERCHEV(code dép. ; $A$4:$C$11 ; 3 ; FAUX) renvoie la zone."
wp["A13"].font = F(italic=True, color=MUTED)

for ref, t in (("E3", "Niveau de tension"), ("F3", "Seuil minimal (candidatures par place)")):
    wp[ref] = t
    head(wp, ref)
for k, (lab, seuil) in enumerate((("Modérée", 10), ("Forte", 20), ("Très forte", 40)), 4):
    wp.cell(k, 5, lab).border = box
    c = wp.cell(k, 6, seuil)
    c.border = box
    c.alignment = Alignment(horizontal="center")
wp["E7"] = "En dessous de 10 : « Faible »."
wp["E7"].font = F(italic=True, color=MUTED)
wp["E9"] = ("Seuils choisis d'après la répartition des 2 691 formations : "
            "environ un quart des formations dans chaque niveau.")
wp["E9"].font = F(italic=True, color=MUTED)
for col, w in (("A", 11), ("B", 20), ("C", 18), ("D", 4), ("E", 20), ("F", 26)):
    wp.column_dimensions[col].width = w
wp.row_dimensions[3].height = 32

# =============================================================================
# Feuille Synthese (formules)
# =============================================================================
wsy = wb.create_sheet("Synthèse")
D = "'Données'!"
rng = lambda col: "%s$%s$2:$%s$%d" % (D, col, col, LAST)

wsy["A1"] = "Tableau de synthèse — candidatures par place selon la filière (Île-de-France, Parcoursup 2025)"
wsy["A1"].font = F(bold=True, size=14, color=NAVY)
wsy["A2"] = "Calculé par formules (NB.SI, SOMME.SI) à partir de la feuille Données. Trié de la filière la plus demandée à la moins demandée."
wsy["A2"].font = F(italic=True, color=MUTED)

H1 = ["Filière", "Formations", "Places", "Candidatures", "Admis",
      "Candidatures par place", "Taux de remplissage", "Part des places"]
for j, t in enumerate(H1, 1):
    wsy.cell(4, j, t)
    head(wsy, "%s4" % get_column_letter(j))
wsy.row_dimensions[4].height = 32

# ordre : pression decroissante (calculee ici uniquement pour ordonner les lignes)
agg = {}
for r in rows:
    f = FILIERE.get(r["fili"], r["fili"])
    a = agg.setdefault(f, [0, 0])
    a[0] += int(r["capa_fin"])
    a[1] += int(r["voe_tot"])
order = sorted(agg, key=lambda f: -agg[f][1] / agg[f][0])
R0 = 5
RT = R0 + len(order)
for k, f in enumerate(order):
    i = R0 + k
    wsy.cell(i, 1, f)
    wsy.cell(i, 2, "=COUNTIF(%s,$A%d)" % (rng("D"), i))
    wsy.cell(i, 3, "=SUMIF(%s,$A%d,%s)" % (rng("D"), i, rng("G")))
    wsy.cell(i, 4, "=SUMIF(%s,$A%d,%s)" % (rng("D"), i, rng("H")))
    wsy.cell(i, 5, "=SUMIF(%s,$A%d,%s)" % (rng("D"), i, rng("I")))
    wsy.cell(i, 6, '=IFERROR(D%d/C%d,"")' % (i, i))
    wsy.cell(i, 7, '=IFERROR(E%d/C%d,"")' % (i, i))
    wsy.cell(i, 8, "=C%d/$C$%d" % (i, RT))
wsy.cell(RT, 1, "Île-de-France")
for j, col in ((2, "B"), (3, "C"), (4, "D"), (5, "E")):
    wsy.cell(RT, j, "=SUM(%s%d:%s%d)" % (col, R0, col, RT - 1))
wsy.cell(RT, 6, "=D%d/C%d" % (RT, RT))
wsy.cell(RT, 7, "=E%d/C%d" % (RT, RT))
wsy.cell(RT, 8, "=SUM(H%d:H%d)" % (R0, RT - 1))


def fmt_block(ws_, r1, r2, ncol, total_row=None):
    for i in range(r1, r2 + 1):
        for j in range(1, ncol + 1):
            c = ws_.cell(i, j)
            c.border = box
            if i == total_row:
                c.font = F(bold=True)
                c.fill = PatternFill("solid", fgColor=PALE)


fmt_block(wsy, R0, RT, 8, RT)
for i in range(R0, RT + 1):
    for j in (2, 3, 4, 5):
        wsy.cell(i, j).number_format = "#,##0"
    wsy.cell(i, 6).number_format = "0.0"
    wsy.cell(i, 7).number_format = "0%"
    wsy.cell(i, 8).number_format = "0.0%"

# --- par zone -----------------------------------------------------------------
Z0 = RT + 3
wsy.cell(Z0 - 1, 1, "Par zone géographique").font = F(bold=True, size=12, color=NAVY)
for j, t in enumerate(["Zone", "Formations", "Places", "Candidatures", "Admis", "Candidatures par place"], 1):
    wsy.cell(Z0, j, t)
    head(wsy, "%s%d" % (get_column_letter(j), Z0))
wsy.row_dimensions[Z0].height = 32
for k, z in enumerate(("Paris", "Petite couronne", "Grande couronne"), 1):
    i = Z0 + k
    wsy.cell(i, 1, z)
    wsy.cell(i, 2, "=COUNTIF(%s,$A%d)" % (rng("J"), i))
    wsy.cell(i, 3, "=SUMIF(%s,$A%d,%s)" % (rng("J"), i, rng("G")))
    wsy.cell(i, 4, "=SUMIF(%s,$A%d,%s)" % (rng("J"), i, rng("H")))
    wsy.cell(i, 5, "=SUMIF(%s,$A%d,%s)" % (rng("J"), i, rng("I")))
    wsy.cell(i, 6, '=IFERROR(D%d/C%d,"")' % (i, i))
fmt_block(wsy, Z0 + 1, Z0 + 3, 6)
for i in range(Z0 + 1, Z0 + 4):
    for j in (2, 3, 4, 5):
        wsy.cell(i, j).number_format = "#,##0"
    wsy.cell(i, 6).number_format = "0.0"

# --- reperes Paris Nanterre -----------------------------------------------------
P0 = Z0 + 7
wsy.cell(P0 - 1, 1, "Repères : mon université").font = F(bold=True, size=12, color=NAVY)
for j, t in enumerate(["Périmètre", "Formations", "Places", "Candidatures", "Admis", "Candidatures par place"], 1):
    wsy.cell(P0, j, t)
    head(wsy, "%s%d" % (get_column_letter(j), P0))
wsy.row_dimensions[P0].height = 32
MIASHS = "Licence - Mathématiques et informatique appliquées aux sciences humaines et sociales"
i = P0 + 1
wsy.cell(i, 1, "Université Paris Nanterre")
wsy.cell(i, 2, '=COUNTIF(%s,"*Paris Nanterre*")' % rng("A"))
for j, col in ((3, "G"), (4, "H"), (5, "I")):
    wsy.cell(i, j, '=SUMIF(%s,"*Paris Nanterre*",%s)' % (rng("A"), rng(col)))
wsy.cell(i, 6, '=IFERROR(D%d/C%d,"")' % (i, i))
i = P0 + 2
wsy.cell(i, 1, "dont licence MIASHS")
wsy.cell(i, 2, '=COUNTIFS(%s,"Université Paris Nanterre",%s,"%s")' % (rng("A"), rng("E"), MIASHS))
for j, col in ((3, "G"), (4, "H"), (5, "I")):
    wsy.cell(i, j, '=SUMIFS(%s,%s,"Université Paris Nanterre",%s,"%s")' % (rng(col), rng("A"), rng("E"), MIASHS))
wsy.cell(i, 6, '=IFERROR(D%d/C%d,"")' % (i, i))
fmt_block(wsy, P0 + 1, P0 + 2, 6)
for i in (P0 + 1, P0 + 2):
    for j in (2, 3, 4, 5):
        wsy.cell(i, j).number_format = "#,##0"
    wsy.cell(i, 6).number_format = "0.0"
wsy.cell(P0 + 3, 1, "Les formations de l'IUT de Ville-d'Avray sont répertoriées sous un autre libellé d'établissement et ne sont pas comptées ici.").font = F(italic=True, color=MUTED)

for col, w in (("A", 26), ("B", 12), ("C", 11), ("D", 14), ("E", 11), ("F", 15), ("G", 14), ("H", 11)):
    wsy.column_dimensions[col].width = w

# =============================================================================
# Feuille Lisez-moi (en premier)
# =============================================================================
wl = wb.create_sheet("Lisez-moi", 0)
wl.sheet_view.showGridLines = False
wl.column_dimensions["A"].width = 3
wl.column_dimensions["B"].width = 28
wl.column_dimensions["C"].width = 72
wl.column_dimensions["D"].width = 24
r = 2
wl.cell(r, 2, "Parcoursup 2025 en Île-de-France").font = F(bold=True, size=20, color=NAVY)
r += 1
wl.cell(r, 2, "Où la concurrence à l'entrée de l'enseignement supérieur est-elle la plus forte ?").font = F(size=13, color=ORANGE, bold=True)
r += 2


def kv(label, value, bold=False):
    global r
    wl.cell(r, 2, label).font = F(bold=True, color=NAVY)
    c = wl.cell(r, 3, value)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    c.font = F(bold=bold)
    wl.cell(r, 2).alignment = Alignment(vertical="top")
    wl.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
    # une cellule fusionnee ne s'ajuste pas seule en hauteur : on l'estime
    wl.row_dimensions[r].height = 15.5 * max(1, -(-len(value) // 112))
    r += 1


def section(title):
    global r
    r += 1
    c = wl.cell(r, 2, title)
    c.font = F(bold=True, size=12, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=NAVY)
    for j in (3, 4):
        wl.cell(r, j).fill = PatternFill("solid", fgColor=NAVY)
    r += 1


kv("Auteur", ("%s — %s" if ID["numero"].startswith("[") else "%s — n° étudiant %s") % (ID["nom_prenom"], ID["numero"]))
kv("Formation", "%s, %s — année universitaire %s" % (ID["filiere"], ID["universite"], ID["annee"]))
kv("Cadre", "EC Compétences numériques : Machines et Logiciels — Séance 5, projet personnel")

section("Problématique")
kv("Question 1", "En Île-de-France, dans quelles filières le nombre de candidatures par place offerte est-il le plus élevé ?")
kv("Question 2", "Cette tension est-elle la même à Paris, en petite couronne et en grande couronne ?")
kv("Indicateur", "Candidatures par place = nombre de candidats ayant confirmé un vœu dans la formation ÷ capacité d'accueil de la formation.")

section("Source des données (données ouvertes)")
kv("Jeu de données", "Parcoursup 2025 — vœux de poursuite d'études et de réorientation dans l'enseignement supérieur et réponses des établissements")
kv("Producteur", "Ministère chargé de l'Enseignement supérieur et de la Recherche — SIES (traitement des données Parcoursup)")
kv("Adresse", "https://data.enseignementsup-recherche.gouv.fr/explore/dataset/fr-esr-parcoursup/")
kv("Licence", "Licence Ouverte v2.0 (Etalab) — réutilisation libre avec mention de la source")
kv("Version", "Jeu mis à jour le 9 mars 2026 ; téléchargé le 12 octobre 2026 au format CSV")
kv("Périmètre retenu", "Filtre Région = Île-de-France : %s formations sur 14 252 (hors apprentissage), 8 départements." % "{:,}".format(N).replace(",", " "))
kv("Colonnes retenues", "9 colonnes sur les 118 du fichier d'origine (voir ci-dessous).")
kv("Nettoyage", "Libellés de filière harmonisés (Licence_Las → Licence L.AS ; Ecole d'Ingénieur → École d'ingénieur ; Ecole de Commerce → École de commerce) et sélectivité abrégée (Sélective / Non sélective). Aucune valeur numérique modifiée.")

section("Feuilles du classeur")
for a, b in (
    ("Données", "Les %s formations. En-têtes bleus : colonnes d'origine. En-têtes orange : colonnes calculées par formule. Tri à deux niveaux (Filière, puis Candidatures par place décroissant) et filtres automatiques." % "{:,}".format(N).replace(",", " ")),
    ("Paramètres", "Table départements → zones (utilisée par RECHERCHEV) et seuils des niveaux de tension (utilisés par les SI imbriqués)."),
    ("Synthèse", "Tableau de synthèse par filière et par zone, calculé par formules (NB.SI, SOMME.SI, SOMME.SI.ENS), et graphique."),
    ("TCD", "Tableaux croisés dynamiques : par filière, puis filière × zone, avec un champ calculé « candidatures par place »."),
    ("Sous-totaux", "Liste triée par filière avec sous-totaux automatiques (Données > Sous-total) des places, candidatures et admis."),
    ("Zoom Paris Nanterre", "Extraction par filtre avancé des formations de mon université, triées de la plus demandée à la moins demandée."),
):
    kv(a, b)

section("Colonnes de la feuille Données")
for a, b, o in (
    ("Établissement", "Nom de l'établissement qui porte la formation", "origine : g_ea_lib_vx"),
    ("Code dép. / Département", "Département de l'établissement", "origine : dep, dep_lib"),
    ("Filière", "Type de formation très agrégé (Licence, BTS, BUT, CPGE, PASS…)", "origine : fili"),
    ("Formation", "Intitulé de la formation", "origine : lib_for_voe_ins"),
    ("Sélectivité", "Formation sélective ou non sélective", "origine : select_form"),
    ("Places", "Capacité de l'établissement pour la formation", "origine : capa_fin"),
    ("Candidatures", "Effectif total des candidats pour la formation", "origine : voe_tot"),
    ("Admis", "Candidats ayant accepté la proposition de l'établissement", "origine : acc_tot"),
    ("Zone", "Paris, Petite couronne ou Grande couronne", "calculée : RECHERCHEV"),
    ("Candidatures par place", "Candidatures ÷ Places", "calculée : SIERREUR"),
    ("Taux de remplissage", "Admis ÷ Places", "calculée : SIERREUR"),
    ("Niveau de tension", "Faible, Modérée, Forte ou Très forte selon les seuils", "calculée : SI imbriqués"),
):
    wl.cell(r, 2, a).font = F(bold=True)
    wl.cell(r, 3, b).alignment = Alignment(wrap_text=True, vertical="top")
    wl.cell(r, 4, o).font = F(color=MUTED, italic=True)
    r += 1

section("Limite de l'indicateur")
kv("À garder en tête", "Un même candidat formule plusieurs vœux : les candidatures s'additionnent d'une formation à l'autre. Le ratio mesure donc l'attractivité d'une formation par rapport à sa capacité, pas la probabilité individuelle d'être admis.")

wb.properties.creator = ID["nom_prenom"]
wb.properties.lastModifiedBy = ID["nom_prenom"]
wb.properties.title = "Parcoursup 2025 en Île-de-France — projet personnel séance 5"
wb.properties.subject = "Projet multimédia — %s" % ID["numero"]

wb.save(OUT)
print("OK", OUT, "| formations :", N, "| ordre synthese :", order)
