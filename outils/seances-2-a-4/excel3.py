# -*- coding: utf-8 -*-
"""Remplit le classeur de la SEANCE 3 (exercices 1 a 6).

Rappel important : dans le format .xlsx les formules sont stockees avec les
noms de fonctions ANGLAIS et des virgules comme separateurs d'arguments.
Excel en francais affichera SOMME, RECHERCHEV, des points-virgules, etc.
"""
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

SRC = r"C:\Users\utilisateur\Desktop\Cours\Environnement Numerique\seance-3\supports\Support_Exercices_Excel_ML_Nanterre.xlsx"
DST = r"C:\Users\utilisateur\Desktop\Cours\Environnement Numerique\seance-3\NOM_Prenom_TD_Excel.xlsx"

EURO = '#,##0.00\\ "€"'
PCT = "0 %"
DATE = "DD/MM/YYYY"

thin = Side(style="thin", color="8EA9DB")
BORD = Border(left=thin, right=thin, top=thin, bottom=thin)
JAUNE = PatternFill("solid", fgColor="FFF2CC")
GRAS = Font(bold=True)

wb = openpyxl.load_workbook(SRC)


def remplir(ws, coord, formule, fmt=None, gras=False):
    c = ws[coord]
    c.value = formule
    c.fill = JAUNE
    c.border = BORD
    if fmt:
        c.number_format = fmt
    if gras:
        c.font = GRAS
    return c


def encadrer(ws, plage):
    for row in ws[plage]:
        for c in row:
            c.border = BORD


def largeurs(ws, mapping):
    for col, w in mapping.items():
        ws.column_dimensions[col].width = w


# =====================================================================
# EXERCICE 1 - releve de depenses
# =====================================================================
ws = wb["Exercice 1"]
for r in range(5, 15):
    ws["A%d" % r].number_format = DATE
    ws["D%d" % r].number_format = EURO
    remplir(ws, "E%d" % r, "=D%d*1.2" % r, EURO)
remplir(ws, "D15", "=SUM(D5:D14)", EURO, gras=True)
remplir(ws, "E15", "=SUM(E5:E14)", EURO, gras=True)
ws["C15"].font = GRAS
encadrer(ws, "A4:E15")
for c in ws["A4:E4"][0]:
    c.font = Font(bold=True, color="FFFFFF")
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
largeurs(ws, {"A": 13, "B": 16, "C": 28, "D": 16, "E": 16})

# =====================================================================
# EXERCICE 2 - references relative / absolue / mixte
# =====================================================================
ws = wb["Exercice 2"]
# Partie A
for r in range(6, 9):
    remplir(ws, "H%d" % r, "=SUM(B%d:G%d)" % (r, r), gras=True)
for col in "BCDEFG":
    remplir(ws, "%s9" % col, "=SUM(%s6:%s8)" % (col, col), gras=True)
remplir(ws, "H9", "=SUM(H6:H8)", gras=True)
encadrer(ws, "A5:H9")

# Partie B
ws["B12"].number_format = PCT
for r in range(15, 18):
    ws["B%d" % r].number_format = EURO
    remplir(ws, "D%d" % r, "=B%d*C%d" % (r, r), EURO)
    remplir(ws, "E%d" % r, "=D%d*(1+$B$12)" % r, EURO)
encadrer(ws, "A14:E17")
encadrer(ws, "A12:B12")

# Partie C - reference mixte : $B21 (colonne bloquee) et C$20 (ligne bloquee)
for col in "CDE":
    ws["%s20" % col].number_format = PCT
for r in range(21, 24):
    ws["B%d" % r].number_format = EURO
    for col in "CDE":
        remplir(ws, "%s%d" % (col, r), "=$B%d*(1-%s$20)" % (r, col), EURO)
encadrer(ws, "A20:E23")
for plage in ("A5:H5", "A14:E14", "A20:E20"):
    for c in ws[plage][0]:
        c.font = Font(bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
largeurs(ws, {"A": 20, "B": 17, "C": 16, "D": 16, "E": 16, "F": 11, "G": 11, "H": 15})

# =====================================================================
# EXERCICE 3 - fonctions logiques
# =====================================================================
ws = wb["Exercice 3"]
for r in range(5, 13):
    remplir(ws, "D%d" % r, "=AVERAGE(B%d:C%d)" % (r, r), "0.00")
    remplir(ws, "E%d" % r, '=IF(D%d>=10,"Admis","Ajourné")' % r)
    remplir(ws, "F%d" % r, '=IF(AND(B%d>=10,C%d>=10),"Oui","Non")' % (r, r))
    remplir(ws, "G%d" % r, '=IF(OR(B%d>=10,C%d>=10),"Oui","Non")' % (r, r))
remplir(ws, "D14", "=COUNTA(A5:A12)", "0", gras=True)
remplir(ws, "D15", "=COUNTA(C5:C12)", "0", gras=True)
encadrer(ws, "A4:G12")
encadrer(ws, "A14:D15")
for c in ws["A4:G4"][0]:
    c.font = Font(bold=True, color="FFFFFF")
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
largeurs(ws, {"A": 14, "B": 14, "C": 14, "D": 11, "E": 15, "F": 17, "G": 22})

# =====================================================================
# EXERCICE 4 - RECHERCHEV, SOMMEPROD, impression
# =====================================================================
ws = wb["Exercice 4"]
for r in range(6, 12):
    ws["C%d" % r].number_format = EURO
for r in range(15, 20):
    remplir(ws, "B%d" % r, "=VLOOKUP(A%d,$A$6:$C$11,2,FALSE)" % r)
    remplir(ws, "C%d" % r, "=VLOOKUP(A%d,$A$6:$C$11,3,FALSE)" % r, EURO)
    remplir(ws, "E%d" % r, "=C%d*D%d" % (r, r), EURO)
remplir(ws, "E20", "=SUM(E15:E19)", EURO, gras=True)
remplir(ws, "E21", "=SUMPRODUCT(C15:C19,D15:D19)", EURO, gras=True)
encadrer(ws, "A5:C11")
encadrer(ws, "A14:E21")
for plage in ("A5:C5", "A14:E14"):
    for c in ws[plage][0]:
        c.font = Font(bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
ws["A4"].font = GRAS
ws["A13"].font = GRAS
ws["A20"].font = GRAS
ws["A21"].font = GRAS
largeurs(ws, {"A": 14, "B": 30, "C": 22, "D": 12, "E": 18})
# mise en page / impression
ws.print_area = "A4:E21"
ws.page_setup.orientation = "portrait"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 1
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.oddHeader.left.text = "[NOM Prénom] - [N° étudiant]"
ws.oddHeader.right.text = "Bon de commande"
ws.oddFooter.center.text = "Page &P / &N"

# =====================================================================
# EXERCICE 5 - gestion de stock
# =====================================================================
ws = wb["Exercice 5"]
for r in range(6, 12):
    ws["B%d" % r].number_format = EURO
for r in range(15, 21):
    remplir(ws, "E%d" % r, "=VLOOKUP(A%d,$A$6:$B$11,2,FALSE)" % r, EURO)
    remplir(ws, "F%d" % r, "=C%d*E%d" % (r, r), EURO)
    remplir(ws, "G%d" % r, '=IF(C%d=0,"Rupture",IF(C%d<=D%d,"Alerte","OK"))' % (r, r, r))
remplir(ws, "F22", "=SUM(F15:F20)", EURO, gras=True)
remplir(ws, "F23", "=AVERAGE(F15:F20)", EURO, gras=True)
remplir(ws, "F24", "=MIN(C15:C20)", "0", gras=True)
remplir(ws, "F25", "=MAX(C15:C20)", "0", gras=True)
encadrer(ws, "A5:B11")
encadrer(ws, "A14:G20")
encadrer(ws, "A22:F25")
for plage in ("A5:B5", "A14:G14"):
    for c in ws[plage][0]:
        c.font = Font(bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
for coord in ("A4", "A13", "A22", "A23", "A24", "A25"):
    ws[coord].font = GRAS
largeurs(ws, {"A": 14, "B": 22, "C": 17, "D": 15, "E": 20, "F": 18, "G": 12})
ws.print_area = "A4:G25"
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 1
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.oddHeader.left.text = "[NOM Prénom] - [N° étudiant]"
ws.oddHeader.right.text = "&D"

# =====================================================================
# EXERCICE 6 - bulletin de notes
# =====================================================================
ws = wb["Exercice 6"]
for r in range(5, 10):
    remplir(ws, "D%d" % r, "=B%d*C%d" % (r, r), "0.00")
remplir(ws, "B10", "=SUM(B5:B9)", "0", gras=True)
remplir(ws, "D10", "=SUM(D5:D9)", "0.00", gras=True)
remplir(ws, "D11", "=D10/B10", "0.00", gras=True)
remplir(ws, "D12", "=VLOOKUP(D11,$F$3:$G$7,2,TRUE)", gras=True)
remplir(ws, "D13", '=IF(AND(D11>=8,D11<10),"Passage devant jury",IF(D11>=10,"Admis","Ajourné"))', gras=True)
remplir(ws, "D14", "=COUNTA(C5:C9)", "0", gras=True)
# reponse au point de reflexion : moyenne calculee sur les seules matieres notees
ws["A15"] = "Moyenne corrigée (matières réellement notées)"
ws["A15"].font = Font(bold=True, italic=True)
remplir(ws, "D15",
        '=SUMPRODUCT(B5:B9,C5:C9)/SUMPRODUCT((C5:C9<>"")*B5:B9)',
        "0.00", gras=True)
encadrer(ws, "A4:D15")
encadrer(ws, "F2:G7")
for plage in ("A4:D4", "F2:G2"):
    for c in ws[plage][0]:
        c.font = Font(bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
for r in range(10, 15):
    ws["A%d" % r].font = GRAS
largeurs(ws, {"A": 34, "B": 13, "C": 11, "D": 26, "E": 3, "F": 10, "G": 14})
ws.print_area = "A1:G15"
ws.page_setup.orientation = "portrait"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 1
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.oddHeader.left.text = "[NOM Prénom] - [N° étudiant]"
ws.oddHeader.right.text = "&D"

# =====================================================================
wb.properties.creator = "[NOM Prénom]"
wb.properties.lastModifiedBy = "[NOM Prénom]"
os.makedirs(os.path.dirname(DST), exist_ok=True)
wb.save(DST)
print("Classeur ecrit :", DST, os.path.getsize(DST), "octets")
